#!/usr/bin/env python3
import os
from docx import Document

source_dir = "Zombie Game"
output_dir = "extracted_docs"

os.makedirs(output_dir, exist_ok=True)

docx_files = [
    "#L01f4d8 GM Companion Guide #U2013 Volume 1_ Core Stat & Trait Companion (Master Edition).docx",
    "Character Creation Quiz.docx",
    "#L01f4d8GM Companion Guide #U2013 Volume 2_ Trait & Flaw Library (Master Edition).docx",
    "Narrative_Progression_System.docx",
    "REVISED_FOUR_SEASONS_STORY_SPINE.docx",
    "#L01f3b2 CORE SYSTEM MODEL_ Tension-Based Minimalism.docx",
    "#L01f4d8 GM Companion Guide_ Volume 4 #U2013 Narrative Hooks, Arc Generators, Red Line Builders.docx",
    "#L01f4d8GM Companion Guide #U2013 Volume 3_ Crisis Profiles & Behavior Engines.docx",
    "#L01f3ad GM Companion Insert_ Crisis Response Layer (Volume 2 Supplement).docx",
    "Action_Cycle_and_Risk_Resolution_System_Clean.docx",
    "#L01f3af RED LINE BUILDER_ GM REFERENCE SHEET.docx",
    "#L01f4d8 GM Companion Guide #U2013 Volume 5_ Resource & Inventory Sysem.docx",
    "GM VOLUME 6_ RESOURCE & INVENTORY SYSTEM.docx",
    "GAME_SYSTEM_INDEX_MASTER.docx"
]

for docx_file in docx_files:
    try:
        doc_path = os.path.join(source_dir, docx_file)
        doc = Document(doc_path)

        # Extract text
        text_content = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_content.append(para.text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    row_text.append(cell.text.strip())
                if any(row_text):
                    text_content.append(" | ".join(row_text))

        # Save to text file
        output_name = docx_file.replace(".docx", ".txt").replace("#L01f4d8", "").replace("#L01f3b2", "").replace("#L01f3ad", "").replace("#L01f3af", "").replace("#L01f4c4", "").replace("#U2013", "-").strip()
        output_path = os.path.join(output_dir, output_name)

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n\n'.join(text_content))

        print(f"✓ Extracted: {output_name}")
    except Exception as e:
        print(f"✗ Error with {docx_file}: {e}")
